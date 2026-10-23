"""Export the organizer workbook into validated, UTF-8 runtime JSON.

The workbook is read only. Runtime modules consume the generated JSON and need
no third-party package.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "Track 3 Data.xlsx"
DATA_DIR = ROOT / "data"

# Sheet name, exact header cells, and canonical output keys. The two leading
# guide sheets are deliberately excluded.
SCHEMAS = {
    "baseline": ("01_机场月度运行基线", ("年份", "月份", "机场名称", "机场分类", "按计划完成航班的比例", "机场放行正常比例", "同类机场平均比例", "该机场过去12个月平均比例", "与自身平均差（百分点）", "样本说明", "来源编号", "数据状态"), ("year", "month", "airport", "tier", "completion", "release", "peer_avg", "own_avg12", "delta_pp", "sample_note", "source", "status")),
    "airports_all": ("02_机场年度规模", ("年份", "机场名称", "城市", "省份", "机场分类", "旅客吞吐量（万人次）", "飞机起降架次", "航空公司业务重要程度", "来源编号", "数据状态"), ("year", "name", "city", "province", "tier", "pax_10k", "movements", "importance", "source", "status")),
    "weather": ("03_机场天气月度", ("年份", "月份", "机场名称", "气象观测站编号", "降水量（毫米）", "平均气温（摄氏度）", "最大风速（米/秒）", "明显不利天气天数", "是否缺少数据", "来源编号", "数据状态"), ("year", "month", "airport", "station", "precipitation_mm", "temp_c", "max_wind_ms", "bad_weather_days", "missing", "source", "status")),
    "events": ("04_模拟事件", ("事件编号", "事件名称", "事件时间", "机场名称", "天气描述", "运行情况描述", "航空公司业务重要程度", "现场备注", "已知信息", "未知信息", "事件状态", "数据标记"), ("id", "name", "time_text", "airport", "weather_text", "ops_text", "importance", "notes", "known", "unknown", "status", "tag")),
    "rules": ("05_运营规则", ("规则编号", "规则主题", "触发条件", "建议核查事项", "建议行动", "建议升级对象", "必须人工审批", "禁止自动执行事项", "规则来源", "规则状态"), ("id", "topic", "trigger", "checks", "action", "escalate_to", "manual_approval", "forbidden", "source", "status")),
    "feedback_seed": ("07_人工反馈", ("事件编号", "建议编号", "人工处理结果", "修改或拒绝原因", "最终采取的行动", "事后结果", "下次改进建议", "复核人", "复核日期"), ("event_id", "rec_id", "result", "reason", "final_action", "outcome", "improvement", "reviewer", "date")),
    "tests": ("08_测试问题", ("测试编号", "测试类型", "输入情境或问题", "预期行为", "禁止行为", "事实核查点", "评分要点", "测试结果", "错误说明", "修正措施"), ("id", "type", "question", "expected", "forbidden", "fact_check", "scoring", "result", "error_note", "fix_note")),
    "field_defs": ("09_字段说明", ("数据表", "字段名称", "中文名称", "含义说明", "数据类型", "单位", "记录层级", "是否允许为空", "允许内容", "来源或计算方法", "质量检查规则", "备注"), ("table", "field", "name_zh", "meaning", "type", "unit", "grain", "nullable", "allowed", "method", "quality_rule", "notes")),
    "sources_list": ("10_数据来源登记", ("来源编号", "发布机构", "资料名称", "网页链接", "发布日期", "数据所属期间", "下载日期", "来源类型", "整理方法", "第一复核人", "第二复核人", "复核状态", "备注"), ("id", "publisher", "title", "url", "published", "period", "downloaded", "type", "method", "reviewer1", "reviewer2", "status", "notes")),
    "data_quality": ("11_数据质量问题", ("问题编号", "数据表", "字段或记录", "问题类型", "问题描述", "可能影响", "处理方式", "负责人", "状态", "复核日期"), ("id", "table", "field", "type", "description", "impact", "handling", "owner", "status", "review_date")),
    "checklist": ("12_提交材料检查表", ("材料编号", "提交材料", "是否必须", "负责人", "状态", "证明材料链接", "复核意见"), ("id", "deliverable", "required", "owner", "status", "evidence", "review")),
}
COUNTS = {"baseline": 192, "weather": 192, "airports_all": 16, "events": 5, "rules": 12, "feedback_seed": 3, "tests": 15, "data_quality": 4, "checklist": 8}


def export(workbook: Path = WORKBOOK) -> dict:
    if not workbook.is_file():
        raise FileNotFoundError(workbook)
    book = load_workbook(workbook, read_only=True, data_only=True)
    tables = {}
    for key, (sheet_name, headers, fields) in SCHEMAS.items():
        if sheet_name not in book.sheetnames:
            raise ValueError(f"Required sheet missing: {sheet_name}")
        rows = book[sheet_name].iter_rows(values_only=True)
        found = tuple(next(rows))
        if found != headers:
            raise ValueError(f"Columns changed in {sheet_name}: {found!r}")
        tables[key] = [dict(zip(fields, row)) for row in rows if any(v is not None for v in row)]
        if key in COUNTS and len(tables[key]) != COUNTS[key]:
            raise ValueError(f"{key}: expected {COUNTS[key]} rows, got {len(tables[key])}")
    book.close()
    for row in tables["weather"]:
        if row["missing"] not in ("是", "否"):
            raise ValueError("Unexpected missing-weather flag")
        row["missing"] = row["missing"] == "是"
    for row in tables["rules"]:
        if row["manual_approval"] not in ("是", "否"):
            raise ValueError("Unexpected manual-approval flag")
        row["manual_approval"] = row["manual_approval"] == "是"
    airports = [a for a in tables.pop("airports_all") if a["year"] == 2025]
    airports.sort(key=lambda a: -a["pax_10k"])
    if len(airports) != 8:
        raise ValueError("Expected eight 2025 airport records")
    tables["airports"] = airports
    tables["airport_names"] = [a["name"] for a in airports]
    tables["sources"] = {r["id"]: r for r in tables.pop("sources_list")}
    tables["meta"] = {"sim_label": "模拟 / SIMULATED", "data_path": str(workbook.name), "counts": {k: len(v) for k, v in tables.items() if isinstance(v, list)}}
    DATA_DIR.mkdir(exist_ok=True)
    (DATA_DIR / "track3.json").write_text(json.dumps(tables, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    (DATA_DIR / "feedback_seed.json").write_text(json.dumps(tables["feedback_seed"], ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    return tables


def write_numbers() -> Path:
    """Render the canonical numbers directly from detector output."""
    import sys

    sys.path.insert(0, str(ROOT))
    from engine.data import load_data
    from engine.detector import ranking

    result = ranking(load_data())
    rows = result["rows"]
    labels = {"海陵国际机场": "Hailing", "云川国际机场": "Yunchuan", "北岭机场": "Beiling"}
    examples = [("海陵国际机场", 2025, 7), ("海陵国际机场", 2025, 8),
                ("云川国际机场", 2025, 7), ("北岭机场", 2025, 12)]
    lines = ["# Canonical detector numbers — 模拟 / SIMULATED", "",
             "Generated by `python scripts/build_data.py` from the organizer workbook and detector output.", "",
             f'- Airport-months: {result["total"]}',
             f'- Own-history flags: {result["flags_own"]}; any-signal flags: {result["flags_any"]}',
             f'- Signal counts: {result["counts"]}',
             "", "| Airport | Month | Own pp | Peer pp | YoY pp | Signals |", "|---|---|---:|---:|---:|---:|"]
    for airport, year, month in examples:
        row = next(r for r in rows if (r["airport"], r["year"], r["month"]) == (airport, year, month))
        lines.append(f'| {labels[airport]} | {row["month_label"]} | {row["delta_pp"]:+.2f} | {row["peer_delta_pp"]:+.2f} | {row["yoy_delta_pp"]:+.2f} | {row["signal_count"]} |')
    lines += ["", "## Network mean completion by calendar month", "",
              "| " + " | ".join(result["monthly_avg_completion"]) + " |",
              "|" + "---:|" * 12,
              "| " + " | ".join(f'{v:.2f}%' for v in result["monthly_avg_completion"].values()) + " |",
              "", "Column I is already in percentage points. Peer and YoY comparisons use completion, not release ratio.",
              "Weather is missing for four rows; those values are not imputed."]
    target = ROOT / "docs" / "NUMBERS.md"
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return target


if __name__ == "__main__":
    result = export()
    write_numbers()
    print("Exported validated workbook:", result["meta"]["counts"])
